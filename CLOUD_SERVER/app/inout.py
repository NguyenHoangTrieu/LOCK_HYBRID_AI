"""
inout.py - Check-in / Check-out tracking, avatar storage, and multi-day Excel (.xlsx) export.
"""
import io
import os
import sqlite3
from datetime import datetime, date, timedelta, timezone
from pathlib import Path
from PIL import Image

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

AVATARS_DIR = Path(__file__).resolve().parent.parent / "data" / "avatars"
AVATARS_DIR.mkdir(parents=True, exist_ok=True)


def save_avatar_image(user_id: int, image_bytes: bytes) -> bool:
    """Saves or updates user profile avatar as JPEG (max 400x400)."""
    try:
        img = Image.open(io.BytesIO(image_bytes))
        if img.mode in ("RGBA", "P"):
            img = img.convert("RGB")
        img.thumbnail((400, 400))
        out_path = AVATARS_DIR / f"{user_id}.jpg"
        img.save(out_path, format="JPEG", quality=85)
        return True
    except Exception as e:
        print(f"Failed to save avatar for user {user_id}: {e}")
        return False


def get_avatar_file(user_id: int) -> Path | None:
    path = AVATARS_DIR / f"{user_id}.jpg"
    return path if path.is_file() else None


def delete_avatar_file(user_id: int) -> None:
    path = AVATARS_DIR / f"{user_id}.jpg"
    if path.is_file():
        try:
            path.unlink()
        except Exception:
            pass


def generate_default_avatar_svg(name: str, user_id: int) -> bytes:
    """Generates a pleasant SVG avatar with the user's initial."""
    initial = (name[0].upper() if name else "U")
    colors = ["#2563EB", "#059669", "#D97706", "#7C3AED", "#DB2777", "#4F46E5", "#0891B2"]
    bg_color = colors[user_id % len(colors)]
    svg = f"""<svg xmlns="http://www.w3.org/2000/svg" width="100" height="100" viewBox="0 0 100 100">
  <circle cx="50" cy="50" r="50" fill="{bg_color}" />
  <text x="50" y="55" font-family="-apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif" font-size="44" font-weight="600" fill="#FFFFFF" text-anchor="middle" dominant-baseline="middle">{initial}</text>
</svg>"""
    return svg.encode("utf-8")


def _to_local_dt(iso_str: str | None) -> datetime | None:
    if not iso_str:
        return None
    try:
        dt = datetime.fromisoformat(iso_str)
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.astimezone()  # Local timezone (+07:00 in Vietnam)
    except Exception:
        return None


def _format_time_only(iso_str: str | None) -> str:
    dt = _to_local_dt(iso_str)
    if dt is not None:
        return dt.strftime("%H:%M:%S")
    return iso_str or ""


def _format_full_datetime(iso_str: str | None) -> str:
    dt = _to_local_dt(iso_str)
    if dt is not None:
        return dt.strftime("%Y-%m-%d %H:%M:%S")
    return (iso_str or "").replace("T", " ").split("+")[0].split(".")[0]


def _format_duration(minutes: int) -> str:
    if minutes <= 0:
        return "0m"
    hours = minutes // 60
    mins = minutes % 60
    if hours > 0:
        return f"{hours}h {mins}m"
    return f"{mins}m"


def get_inout_status_for_date(db: sqlite3.Connection, target_date: str) -> dict:
    """
    Returns daily check-in / check-out status for all active users on `target_date` (YYYY-MM-DD).
    Queries events based on the local day's equivalent UTC range.
    """
    users = db.execute(
        "SELECT id, name, org_id FROM users WHERE active = 1 ORDER BY name ASC"
    ).fetchall()

    user_list = []
    currently_in = 0
    currently_out = 0
    not_checked_in = 0

    try:
        dt0 = datetime.strptime(target_date, "%Y-%m-%d").astimezone()
        start_utc_iso = dt0.replace(hour=0, minute=0, second=0, microsecond=0).astimezone(timezone.utc).isoformat()
        end_utc_iso = dt0.replace(hour=23, minute=59, second=59, microsecond=999999).astimezone(timezone.utc).isoformat()
    except Exception:
        start_utc_iso = f"{target_date}T00:00:00"
        end_utc_iso = f"{target_date}T23:59:59"

    now_utc = datetime.now(timezone.utc)
    is_today = (target_date == datetime.now().strftime("%Y-%m-%d"))

    for u in users:
        uid = u["id"]
        # Fetch all granted events for user on this local date
        events = db.execute(
            """
            SELECT a.id, a.method, a.outcome, a.confidence, a.occurred_at, a.direction, d.name AS device_name
            FROM attendance_events a
            LEFT JOIN devices d ON d.id = a.device_id
            WHERE a.user_id = ? AND a.outcome = 'granted'
              AND a.occurred_at >= ? AND a.occurred_at <= ?
            ORDER BY a.occurred_at ASC, a.id ASC
            """,
            (uid, start_utc_iso, end_utc_iso),
        ).fetchall()

        avatar_exists = (AVATARS_DIR / f"{uid}.jpg").is_file()

        if not events:
            not_checked_in += 1
            user_list.append({
                "user_id": uid,
                "name": u["name"],
                "org_id": u["org_id"] or "",
                "status": "not_checked_in",
                "first_in": None,
                "first_in_time": "",
                "last_out": None,
                "last_out_time": "",
                "last_event_time": "",
                "last_method": None,
                "last_device": None,
                "duration_minutes": 0,
                "duration_str": "0m",
                "scans_count": 0,
                "avatar_url": f"/admin/api/users/{uid}/avatar",
                "has_custom_avatar": avatar_exists,
                "events": [],
            })
            continue

        scans_count = len(events)
        first_event = events[0]
        last_event = events[-1]

        current_status = last_event["direction"] or "in"
        if current_status == "in":
            currently_in += 1
        else:
            currently_out += 1

        first_in = first_event["occurred_at"]
        out_events = [e for e in events if e["direction"] == "out"]
        last_out = out_events[-1]["occurred_at"] if out_events else None

        # Calculate presence duration in minutes
        total_duration_sec = 0.0
        in_time = None

        for ev in events:
            try:
                ev_dt = datetime.fromisoformat(ev["occurred_at"])
                if ev_dt.tzinfo is None:
                    ev_dt = ev_dt.replace(tzinfo=timezone.utc)
            except Exception:
                continue

            if ev["direction"] == "in":
                in_time = ev_dt
            elif ev["direction"] == "out" and in_time is not None:
                total_duration_sec += (ev_dt - in_time).total_seconds()
                in_time = None

        if in_time is not None and is_today:
            # Still checked in today
            total_duration_sec += max(0, (now_utc - in_time).total_seconds())

        duration_mins = int(total_duration_sec // 60)

        user_list.append({
            "user_id": uid,
            "name": u["name"],
            "org_id": u["org_id"] or "",
            "status": current_status,
            "first_in": first_in,
            "first_in_time": _format_time_only(first_in),
            "last_out": last_out,
            "last_out_time": _format_time_only(last_out),
            "last_event_time": _format_time_only(last_event["occurred_at"]),
            "last_method": last_event["method"],
            "last_device": last_event["device_name"] or "Edge Node",
            "duration_minutes": duration_mins,
            "duration_str": _format_duration(duration_mins),
            "scans_count": scans_count,
            "avatar_url": f"/admin/api/users/{uid}/avatar",
            "has_custom_avatar": avatar_exists,
            "events": [
                {
                    "id": e["id"],
                    "time": _format_time_only(e["occurred_at"]),
                    "direction": e["direction"] or "in",
                    "method": e["method"],
                    "device": e["device_name"] or "Edge Node",
                }
                for e in events
            ],
        })

    return {
        "date": target_date,
        "summary": {
            "total_users": len(users),
            "currently_in": currently_in,
            "currently_out": currently_out,
            "not_checked_in": not_checked_in,
        },
        "users": user_list,
    }


def generate_inout_excel(db: sqlite3.Connection, from_date_str: str, to_date_str: str) -> bytes:
    """
    Generates a multi-day Excel report (.xlsx) with:
    - Sheet 1: Daily Summary (Date, User, Org, First In, Last Out, Total Duration, Scans, Status)
    - Sheet 2: Detailed Scan Logs (Timestamp, User, Org, Direction, Method, Device, Confidence)
    """
    wb = openpyxl.Workbook()

    # Style definitions
    font_title = Font(name="Segoe UI", size=16, bold=True, color="1F4E78")
    font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="595959")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_body = Font(name="Segoe UI", size=10)
    font_bold = Font(name="Segoe UI", size=10, bold=True)

    fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    fill_in = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fill_out = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    fill_absent = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    color_in_text = Font(name="Segoe UI", size=10, bold=True, color="276A3C")
    color_out_text = Font(name="Segoe UI", size=10, bold=True, color="C65911")
    color_absent_text = Font(name="Segoe UI", size=10, color="7F7F7F")

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # -------------------------------------------------------------------------
    # Sheet 1: Tổng Hợp Vào-Ra (Daily Summary)
    # -------------------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Tổng Hợp Vào-Ra"
    ws1.views.sheetView[0].showGridLines = True

    # Title & Metadata
    ws1.append(["BÁO CÁO ĐIỂM DANH VÀO / RA"])
    ws1.cell(row=1, column=1).font = font_title
    ws1.append([f"Khoảng thời gian: {from_date_str} đến {to_date_str}  |  Xuất ngày: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
    ws1.cell(row=2, column=1).font = font_subtitle
    ws1.append([])  # blank row

    headers_ws1 = [
        "Ngày", "Mã NV / Org", "Họ và Tên",
        "Giờ Vào Đầu (Check-In)", "Giờ Ra Cuối (Check-Out)",
        "Tổng Thời Gian", "Số Lần Quét", "Trạng Thái"
    ]
    ws1.append(headers_ws1)
    header_row_idx = 4
    for col_idx in range(1, len(headers_ws1) + 1):
        cell = ws1.cell(row=header_row_idx, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center

    # Iterate through all dates in range
    start_d = datetime.strptime(from_date_str, "%Y-%m-%d").date()
    end_d = datetime.strptime(to_date_str, "%Y-%m-%d").date()
    curr_d = start_d

    row_count_ws1 = 4
    while curr_d <= end_d:
        d_str = curr_d.strftime("%Y-%m-%d")
        daily_data = get_inout_status_for_date(db, d_str)

        for u in daily_data["users"]:
            # Only include users who had at least 1 scan or if date range is small
            if u["scans_count"] == 0 and (end_d - start_d).days > 7:
                continue

            row_count_ws1 += 1
            status_text = "CHƯA VÀO"
            status_fill = fill_absent
            status_font = color_absent_text
            if u["status"] == "in":
                status_text = "ĐANG CÓ MẶT (IN)"
                status_fill = fill_in
                status_font = color_in_text
            elif u["status"] == "out":
                status_text = "ĐÃ RA NGOÀI (OUT)"
                status_fill = fill_out
                status_font = color_out_text

            ws1.append([
                d_str,
                u["org_id"],
                u["name"],
                u["first_in_time"] or "--:--:--",
                u["last_out_time"] or "--:--:--",
                u["duration_str"],
                u["scans_count"],
                status_text,
            ])

            # Apply cell styles
            ws1.cell(row=row_count_ws1, column=1).alignment = align_center
            ws1.cell(row=row_count_ws1, column=2).alignment = align_center
            ws1.cell(row=row_count_ws1, column=3).alignment = align_left
            ws1.cell(row=row_count_ws1, column=4).alignment = align_center
            ws1.cell(row=row_count_ws1, column=5).alignment = align_center
            ws1.cell(row=row_count_ws1, column=6).alignment = align_center
            ws1.cell(row=row_count_ws1, column=7).alignment = align_center
            
            c_status = ws1.cell(row=row_count_ws1, column=8)
            c_status.alignment = align_center
            c_status.fill = status_fill
            c_status.font = status_font

            for col_idx in range(1, len(headers_ws1) + 1):
                ws1.cell(row=row_count_ws1, column=col_idx).border = thin_border
                if col_idx != 8:
                    ws1.cell(row=row_count_ws1, column=col_idx).font = font_body

        curr_d += timedelta(days=1)

    # -------------------------------------------------------------------------
    # Sheet 2: Chi Tiết Nhật Ký (Raw Log)
    # -------------------------------------------------------------------------
    ws2 = wb.create_sheet("Nhật Ký Quét Chi Tiết")
    ws2.views.sheetView[0].showGridLines = True

    ws2.append(["NHẬT KÝ QUÉT CHI TIẾT (LOGS)"])
    ws2.cell(row=1, column=1).font = font_title
    ws2.append([f"Từ: {from_date_str} đến: {to_date_str}"])
    ws2.cell(row=2, column=1).font = font_subtitle
    ws2.append([])

    headers_ws2 = [
        "STT", "Thời Gian Quét", "Mã NV / Org", "Họ và Tên",
        "Hướng (Direction)", "Kết Quả", "Phương Thức", "Thiết Bị", "Độ Tin Cậy"
    ]
    ws2.append(headers_ws2)
    for col_idx in range(1, len(headers_ws2) + 1):
        cell = ws2.cell(row=4, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center

    # Query all events in date range (converting local day range to UTC)
    try:
        from_dt = datetime.strptime(from_date_str, "%Y-%m-%d").astimezone()
        start_utc = from_dt.replace(hour=0, minute=0, second=0, microsecond=0).astimezone(timezone.utc).isoformat()
        to_dt = datetime.strptime(to_date_str, "%Y-%m-%d").astimezone()
        end_utc = to_dt.replace(hour=23, minute=59, second=59, microsecond=999999).astimezone(timezone.utc).isoformat()
    except Exception:
        start_utc = f"{from_date_str}T00:00:00"
        end_utc = f"{to_date_str}T23:59:59"

    events = db.execute(
        """
        SELECT a.id, a.user_id, a.method, a.outcome, a.confidence, a.occurred_at, a.direction,
               u.name AS user_name, u.org_id, d.name AS device_name
        FROM attendance_events a
        LEFT JOIN users u ON u.id = a.user_id
        LEFT JOIN devices d ON d.id = a.device_id
        WHERE a.occurred_at >= ? AND a.occurred_at <= ?
        ORDER BY a.occurred_at DESC, a.id DESC
        """,
        (start_utc, end_utc),
    ).fetchall()

    row_count_ws2 = 4
    for idx, ev in enumerate(events, 1):
        row_count_ws2 += 1
        dir_text = "VÀO (CHECK-IN)" if ev["direction"] == "in" else ("RA (CHECK-OUT)" if ev["direction"] == "out" else "--")
        method_text = "Khuôn mặt" if ev["method"] == "face" else "Thẻ RFID"
        outcome_text = "Chấp thuận" if ev["outcome"] == "granted" else "Từ chối"
        conf_text = f"{round(ev['confidence'] * 100, 1)}%" if ev["confidence"] is not None else "--"

        ws2.append([
            idx,
            _format_full_datetime(ev["occurred_at"]),
            ev["org_id"] or "",
            ev["user_name"] or "(Chưa đăng ký)",
            dir_text,
            outcome_text,
            method_text,
            ev["device_name"] or "Edge Node",
            conf_text,
        ])

        ws2.cell(row=row_count_ws2, column=1).alignment = align_center
        ws2.cell(row=row_count_ws2, column=2).alignment = align_center
        ws2.cell(row=row_count_ws2, column=3).alignment = align_center
        ws2.cell(row=row_count_ws2, column=4).alignment = align_left
        ws2.cell(row=row_count_ws2, column=5).alignment = align_center
        ws2.cell(row=row_count_ws2, column=6).alignment = align_center
        ws2.cell(row=row_count_ws2, column=7).alignment = align_center
        ws2.cell(row=row_count_ws2, column=8).alignment = align_left
        ws2.cell(row=row_count_ws2, column=9).alignment = align_center

        if ev["direction"] == "in":
            ws2.cell(row=row_count_ws2, column=5).fill = fill_in
            ws2.cell(row=row_count_ws2, column=5).font = color_in_text
        elif ev["direction"] == "out":
            ws2.cell(row=row_count_ws2, column=5).fill = fill_out
            ws2.cell(row=row_count_ws2, column=5).font = color_out_text

        for col_idx in range(1, len(headers_ws2) + 1):
            ws2.cell(row=row_count_ws2, column=col_idx).border = thin_border
            if col_idx != 5:
                ws2.cell(row=row_count_ws2, column=col_idx).font = font_body

    # Auto-adjust column widths for both sheets
    for ws in [ws1, ws2]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                # Ignore title rows when calculating length
                if cell.row in (1, 2, 3):
                    continue
                max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

    out_buf = io.BytesIO()
    wb.save(out_buf)
    return out_buf.getvalue()
